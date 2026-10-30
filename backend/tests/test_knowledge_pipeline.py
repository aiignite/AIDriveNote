"""知识入库 → 解析 → 蒸馏 → 编译 的核心链路测试。"""
from __future__ import annotations

import io
import uuid
from pathlib import Path

import pytest
from sqlalchemy import func, select

from app.models.knowledge import IngestDocument, IngestJob, KnowledgeCard, NoteSource
from app.services.knowledge.ingest_service import ingest
from app.services.knowledge.ingest_worker import _distill_document, _parse_document
from app.services.knowledge.knowledge_compile_service import build_kb_target, compile_card
from app.services.knowledge.knowledge_extract_service import extract_cards, save_cards
from app.services.knowledge.parsing.base import ParseError
from app.services.knowledge.parsing.registry import parse_bytes, parse_text
from app.services.knowledge.rule_extractor import build_rule_card
from app.services.auth_service import AuthService


@pytest.mark.asyncio
async def test_ingest_is_idempotent_and_enqueues_parse(db_session, test_user):
    first, duplicate_first = await ingest(
        db_session,
        test_user.id,
        source_type="email",
        external_id="message-1",
        title="测试邮件",
        text="第一段内容。\n\n第二段内容。",
    )
    second, duplicate_second = await ingest(
        db_session,
        test_user.id,
        source_type="email",
        external_id="message-1",
        title="测试邮件",
        text="第一段内容。\n\n第二段内容。",
    )

    assert duplicate_first is False
    assert duplicate_second is True
    assert first.id == second.id

    doc_count = await db_session.scalar(select(func.count()).select_from(IngestDocument))
    job_count = await db_session.scalar(select(func.count()).select_from(IngestJob))
    assert doc_count == 1
    assert job_count == 1
    job = (await db_session.execute(select(IngestJob))).scalar_one()
    assert job.stage == "parse"
    assert job.status == "pending"


@pytest.mark.asyncio
async def test_ingest_bytes_saves_temp_file_for_parse(db_session, test_user):
    document, duplicate = await ingest(
        db_session,
        test_user.id,
        source_type="upload",
        external_id="upload-1",
        title="hello.txt",
        data=b"hello knowledge",
        mime="text/plain",
        raw_meta={"filename": "hello.txt"},
    )
    assert duplicate is False
    path = Path(str((document.raw_meta or {})["file_path"]))
    try:
        assert path.exists()
        assert path.read_bytes() == b"hello knowledge"
    finally:
        path.unlink(missing_ok=True)


def test_parse_text_and_html():
    plain = parse_text("第一段\n\n第二段", title="标题")
    assert plain.parser == "text"
    assert "第一段" in plain.text
    assert plain.meta["title"] == "标题"

    html = parse_text(
        "<html><head><title>文章标题</title></head>"
        "<body><h1>小标题</h1><p>这里是正文内容。</p></body></html>",
        mime="text/html",
    )
    assert html.parser == "html"
    assert "正文内容" in html.text


def test_parse_bytes_markdown_and_unknown_binary():
    parsed = parse_bytes("# 标题\n\n正文".encode("utf-8"), filename="a.md")
    assert parsed.parser == "text"
    assert "# 标题" in parsed.text

    with pytest.raises(ParseError):
        parse_bytes(b"\x00\x01\x02\x03\x04\x05", filename="a.bin")


def test_parse_pptx_extracts_slide_text():
    pytest.importorskip("pptx")
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    textbox.text_frame.text = "出风口执行器选型要点"
    buffer = io.BytesIO()
    presentation.save(buffer)

    parsed = parse_bytes(buffer.getvalue(), filename="选型.pptx")
    assert parsed.parser == "pptx"
    assert "Slide 1" in parsed.text
    assert "出风口执行器" in parsed.text


def _make_doc(db_session, test_user):
    return IngestDocument(
        user_id=test_user.id,
        source_type="email",
        external_id="rule-doc",
        content_hash="a" * 64,
        title="执行器选型讨论",
        raw_text=(
            "# 执行器选型讨论\n\n"
            "本次讨论聚焦出风口执行器的供应商与产能。\n\n"
            "## 结论\n\n"
            "- 优先选择博世作为主供应商\n"
            "- 需要提前锁定产能\n\n"
            "## 风险\n\n"
            "- 单一供应商存在交付风险\n"
        ),
        raw_meta={},
        parse_meta={},
        status="parsed",
    )


@pytest.mark.asyncio
async def test_rule_extractor_builds_actionable_card(db_session, test_user):
    doc = _make_doc(db_session, test_user)
    db_session.add(doc)
    await db_session.commit()

    card = build_rule_card(doc)
    assert card.title == "执行器选型讨论"
    assert "供应商" in card.summary
    assert card.payload["key_points"]
    assert card.payload["keywords"]
    assert card.confidence == "low"


@pytest.mark.asyncio
async def test_worker_parse_then_distill_creates_pending_card(db_session, test_user):
    document, _duplicate = await ingest(
        db_session,
        test_user.id,
        source_type="email",
        external_id="pipeline-1",
        title="执行器选型讨论",
        text=_make_doc(db_session, test_user).raw_text or "",
    )

    await _parse_document(db_session, document)
    await db_session.refresh(document)
    assert document.status == "parsed"
    assert document.parser == "text"
    assert "供应商" in (document.raw_text or "")

    await _distill_document(db_session, document)
    await db_session.refresh(document)
    assert document.status == "pending_review"
    cards = (
        await db_session.execute(
            select(KnowledgeCard).where(KnowledgeCard.document_id == document.id)
        )
    ).scalars().all()
    assert cards
    card = cards[0]
    assert card.status == "pending_review"
    assert card.extractor == "rule"
    assert "供应商" in card.summary


@pytest.mark.asyncio
async def test_email_digest_splits_into_multiple_cards(db_session, test_user):
    text = (
        "# 每日摘要\n\n"
        "## 条目一：供应商产能\n\n"
        "博世需要提前锁定产能。\n\n"
        "## 条目二：交付风险\n\n"
        "单一供应商存在交付风险。\n"
    )
    document, _duplicate = await ingest(
        db_session,
        test_user.id,
        source_type="email",
        external_id="digest-1",
        title="每日摘要",
        text=text,
    )
    drafts = await extract_cards(db_session, document, test_user.id)
    assert len(drafts) == 2
    assert drafts[0].title.startswith("每日摘要")
    cards = await save_cards(db_session, document, drafts)
    assert len(cards) == 2
    assert len({card.card_key for card in cards}) == 2


@pytest.mark.asyncio
async def test_compile_card_creates_note_with_source_and_tags(db_session, test_user):
    document, _duplicate = await ingest(
        db_session,
        test_user.id,
        source_type="email",
        external_id="compile-1",
        title="执行器选型讨论",
        text=_make_doc(db_session, test_user).raw_text or "",
    )
    drafts = await extract_cards(db_session, document, test_user.id)
    cards = await save_cards(db_session, document, drafts)
    card = cards[0]

    note, kb_target = await compile_card(db_session, test_user.id, card.id, mode="new")
    assert note.title == card.title
    assert "摘要" in str(note.content.get("text", ""))
    assert "[kb]" in (note.description or "")
    assert kb_target["wikiPath"].startswith("wiki/")
    assert kb_target["rawPath"].startswith("raw/")

    await db_session.refresh(card)
    await db_session.refresh(document)
    assert card.status == "approved"
    assert card.note_id == note.id
    assert document.status == "approved"

    source = (
        await db_session.execute(select(NoteSource).where(NoteSource.note_id == note.id))
    ).scalar_one()
    assert source.document_id == document.id
    assert source.role == "primary"


def test_build_kb_target_uses_expected_dirs():
    doc = IngestDocument(
        source_type="clip",
        external_id="x",
        content_hash="b" * 64,
        title="网页文章",
        raw_meta={},
    )
    card = KnowledgeCard(
        knowledge_type="concept",
        title="出风口执行器",
        summary="一句话摘要",
        confidence="high",
        payload={"topic": "产品开发", "tags": ["执行器"]},
    )
    target = build_kb_target(card, doc)
    assert target["wikiPath"].startswith("wiki/concepts/")
    assert target["rawPath"].startswith("raw/web/")
    assert target["frontmatter"]["type"] == "concept"
    assert target["frontmatter"]["sources"] == [target["rawPath"]]


@pytest.mark.asyncio
async def test_ingest_inbox_approve_api(client, db_session, test_user):
    access_token, _ = AuthService.create_access_token(test_user.id)
    headers = {"Authorization": f"Bearer {access_token}"}

    created = await client.post(
        "/api/v1/ingest/documents",
        json={
            "sourceType": "email",
            "externalId": "api-mail-1",
            "title": "执行器选型讨论",
            "text": (
                "# 执行器选型讨论\n\n"
                "本次讨论聚焦出风口执行器的供应商与产能。\n\n"
                "## 结论\n\n"
                "- 优先选择博世作为主供应商\n"
                "- 需要提前锁定产能\n"
            ),
        },
        headers=headers,
    )
    assert created.status_code == 200, created.text
    document_id = created.json()["documentId"]

    document = await db_session.get(IngestDocument, uuid.UUID(document_id))
    assert document is not None
    await _parse_document(db_session, document)
    await _distill_document(db_session, document)

    inbox = await client.get(
        "/api/v1/knowledge/inbox?status=pending_review", headers=headers,
    )
    assert inbox.status_code == 200, inbox.text
    items = inbox.json()["items"]
    assert len(items) == 1
    card = items[0]
    assert card["title"] == "执行器选型讨论"
    assert card["source"]["type"] == "email"

    approved = await client.post(
        f"/api/v1/knowledge/cards/{card['id']}/approve",
        json={"mode": "new"},
        headers=headers,
    )
    assert approved.status_code == 200, approved.text
    note_id = approved.json()["noteId"]

    note = await client.get(f"/api/v1/notes/{note_id}", headers=headers)
    assert note.status_code == 200, note.text
    body = note.json()
    assert body["title"] == "执行器选型讨论"
    assert "摘要" in str(body.get("content", {}).get("text", ""))

    search = await client.get(
        "/api/v1/notes/search/fulltext?q=执行器", headers=headers,
    )
    assert search.status_code == 200, search.text
    payload = search.json()
    assert payload["knowledge"][note_id]["source_type"] == "email"
    assert payload["knowledge"][note_id]["summary"]

    entity = await client.post(
        f"/api/v1/knowledge/cards/{card['id']}/entity-page",
        json={"name": "博世", "entityType": "entity"},
        headers=headers,
    )
    assert entity.status_code == 200, entity.text
    assert entity.json()["created"] is True
    again = await client.post(
        f"/api/v1/knowledge/cards/{card['id']}/entity-page",
        json={"name": "博世", "entityType": "entity"},
        headers=headers,
    )
    assert again.status_code == 200
    assert again.json()["created"] is False

    lint = await client.get("/api/v1/knowledge/lint", headers=headers)
    assert lint.status_code == 200, lint.text
    assert lint.json()["checked"] >= 2


@pytest.mark.asyncio
async def test_knowledge_settings_api(client, test_user):
    access_token, _ = AuthService.create_access_token(test_user.id)
    headers = {"Authorization": f"Bearer {access_token}"}

    current = await client.get("/api/v1/knowledge/settings", headers=headers)
    assert current.status_code == 200, current.text
    data = current.json()
    assert "distillEnabled" in data
    assert data["dailyLimit"] >= 0

    patched = await client.patch(
        "/api/v1/knowledge/settings",
        json={"allowExternal": True, "dailyLimit": 7},
        headers=headers,
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["allowExternal"] is True
    assert patched.json()["dailyLimit"] == 7


@pytest.mark.asyncio
async def test_knowledge_distill_model_selection(client, test_user):
    access_token, _ = AuthService.create_access_token(test_user.id)
    headers = {"Authorization": f"Bearer {access_token}"}

    current = await client.get("/api/v1/knowledge/settings", headers=headers)
    assert current.status_code == 200, current.text
    assert "defaultModel" in current.json()

    patched = await client.patch(
        "/api/v1/knowledge/settings",
        json={"model": "qwen-max"},
        headers=headers,
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["model"] == "qwen-max"

    cleared = await client.patch(
        "/api/v1/knowledge/settings",
        json={"model": ""},
        headers=headers,
    )
    assert cleared.status_code == 200, cleared.text
    assert cleared.json()["model"] in (None, "")


@pytest.mark.asyncio
async def test_failed_document_is_reenqueued_on_resubmit(db_session, test_user):
    """失败文档不会产出卡片，重复提交应视为重试而不是静默去重。"""
    from sqlalchemy import update

    first, duplicate_first = await ingest(
        db_session,
        test_user.id,
        source_type="upload",
        external_id="retry-1",
        title="失败文档",
        text="重新提交后应再次入队。",
    )
    assert duplicate_first is False
    await db_session.execute(
        update(IngestDocument)
        .where(IngestDocument.id == first.id)
        .values(status="failed", error="解析失败")
    )
    await db_session.commit()

    retried, duplicate_retried = await ingest(
        db_session,
        test_user.id,
        source_type="upload",
        external_id="retry-1",
        title="失败文档",
        text="重新提交后应再次入队。",
    )

    assert duplicate_retried is False
    assert retried.id == first.id
    assert retried.status == "received"
    assert retried.error is None
    job = (await db_session.execute(select(IngestJob))).scalar_one()
    assert job.document_id == first.id
    assert job.stage == "parse"
    assert job.status == "pending"


@pytest.mark.asyncio
async def test_distill_notes_enqueues_tagged_and_document_notes(db_session, test_user, client):
    """批量蒸馏覆盖 kb: 笔记与已上传文档，并与单篇抽取共用幂等键。"""
    from app.models.note.document import NoteDocument
    from app.services.note.note_enhance_service import NoteTagService
    from app.services.note.note_service import NoteService

    kb_note = await NoteService.create_note(db_session, {
        "title": "执行器选型",
        "note_type": "markdown",
        "content": {"text": "# 执行器选型\n\n供应商产能需要提前锁定。"},
        "created_by": test_user.id,
        "updated_by": test_user.id,
    })
    await NoteService.create_note(db_session, {
        "title": "普通笔记",
        "note_type": "markdown",
        "content": {"text": "不应被批量蒸馏"},
        "created_by": test_user.id,
        "updated_by": test_user.id,
    })
    doc_note = await NoteService.create_note(db_session, {
        "title": "上传的标准",
        "note_type": "document",
        "content": None,
        "created_by": test_user.id,
        "updated_by": test_user.id,
    })
    db_session.add(NoteDocument(
        note_id=doc_note.id,
        file_name="standard.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ext=".docx",
        size_bytes=10,
        sha256="a" * 64,
        storage_path="documents/standard.docx",
        etag="a" * 64,
        meta={},
        extract_status="ok",
        extracted_text="设备噪声不得超过 75 分贝。",
    ))
    await db_session.commit()
    tag = await NoteTagService.create_tag(db_session, test_user.id, "kb:concept")
    await NoteTagService.add_tag_to_note(db_session, kb_note.id, tag.id, test_user.id)

    access_token, _ = AuthService.create_access_token(test_user.id)
    res = await client.post(
        "/api/v1/knowledge/distill/notes",
        headers={"Authorization": f"Bearer {access_token}"},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["total"] == 2
    assert body["queued"] == 2
    assert body["duplicate"] == 0
    assert body["skipped"] == 0

    documents = (
        await db_session.execute(select(IngestDocument).order_by(IngestDocument.source_type))
    ).scalars().all()
    assert [d.source_type for d in documents] == ["note", "upload"]
    by_type = {d.source_type: d for d in documents}
    assert by_type["note"].external_id == f"note:{kb_note.id}"
    assert by_type["upload"].external_id == f"note-doc:{'a' * 64}"
    assert by_type["upload"].document_note_id == doc_note.id

    again = await client.post(
        "/api/v1/knowledge/distill/notes",
        headers={"Authorization": f"Bearer {access_token}"},
    )
    assert again.json()["duplicate"] == 2
    assert again.json()["queued"] == 0


@pytest.mark.asyncio
async def test_call_llm_really_calls_model(monkeypatch, db_session, test_user):
    """LLM 抽取必须真的发起模型调用。

    单测环境会跳过模型调用，超时参数一旦取不到就会整条链路退回规则抽取
    （卡片全是低置信度），因此这里用假 provider 单独覆盖这一步。
    """
    from types import SimpleNamespace

    from app.services.knowledge import knowledge_extract_service as svc

    monkeypatch.delenv("AIDRIVE_TESTING", raising=False)

    class _FakeProvider:
        async def chat_with_tools(self, messages, options):
            return {"content": '{"cards": []}'}

    async def _fake_resolve(db, user_id, **kwargs):
        return SimpleNamespace(
            provider="OLLAMA",
            model_id="qwen2.5",
            endpoint="http://ollama:11434",
            api_key="",
            model_name="Ollama/qwen2.5",
        )

    monkeypatch.setattr(svc.AIProviderFactory, "create", lambda *args, **kwargs: _FakeProvider())
    monkeypatch.setattr(svc.LLMRouter, "resolve", _fake_resolve)

    result = await svc._call_llm(db_session, test_user.id, "正文")
    assert result is not None
    assert result[0] == '{"cards": []}'


@pytest.mark.asyncio
async def test_retry_all_cards_requeues_documents(db_session, test_user, client):
    """「用模型重跑」把待审卡片作废，并让来源文档重新入队蒸馏。"""
    document, _duplicate = await ingest(
        db_session,
        test_user.id,
        source_type="note",
        external_id="retry-all-1",
        title="待重跑文档",
        text="正文内容。",
    )
    card = KnowledgeCard(
        user_id=test_user.id,
        document_id=document.id,
        version=1,
        card_key="card-1",
        status="pending_review",
        knowledge_type="concept",
        title="标题",
        summary="摘要",
        confidence="low",
        payload={},
        extractor="rule",
    )
    db_session.add(card)
    await db_session.commit()

    access_token, _ = AuthService.create_access_token(test_user.id)
    res = await client.post(
        "/api/v1/knowledge/cards/retry-all",
        headers={"Authorization": f"Bearer {access_token}"},
    )
    assert res.status_code == 200, res.text
    assert res.json()["documents"] == 1

    await db_session.refresh(card)
    await db_session.refresh(document)
    assert card.status == "ignored"
    assert document.status == "parsed"
    job = (
        await db_session.execute(select(IngestJob).where(IngestJob.stage == "distill"))
    ).scalar_one()
    assert job.document_id == document.id


@pytest.mark.asyncio
async def test_search_knowledge_only_returns_library_cards(db_session, test_user, client):
    """知识卡检索只返回已入库的卡，并同时被 AI 工具 / HTTP 接口复用。"""
    from app.ai_tools.registry import ToolRegistry
    from app.services.knowledge.search_service import search_cards

    document, _duplicate = await ingest(
        db_session,
        test_user.id,
        source_type="upload",
        external_id="search-1",
        title="采购标准",
        text="设备噪声不得超过 75 分贝。",
    )
    approved = KnowledgeCard(
        user_id=test_user.id,
        document_id=document.id,
        version=1,
        card_key="card-1",
        status="approved",
        knowledge_type="decision",
        title="噪声限值决策",
        summary="设备噪声不得超过 75 分贝",
        confidence="high",
        payload={"key_points": ["75 分贝"], "keywords": ["噪声"]},
        extractor="llm",
    )
    pending = KnowledgeCard(
        user_id=test_user.id,
        document_id=document.id,
        version=1,
        card_key="card-2",
        status="pending_review",
        knowledge_type="concept",
        title="未审核噪声卡",
        summary="待审核内容",
        confidence="low",
        payload={},
        extractor="rule",
    )
    db_session.add_all([approved, pending])
    await db_session.commit()

    cards = await search_cards(db_session, test_user.id, "噪声")
    assert [card["title"] for card in cards] == ["噪声限值决策"]

    tool_result = await ToolRegistry.execute(
        "search_knowledge", db_session, test_user.id, query="噪声",
    )
    assert tool_result["success"] is True
    assert tool_result["found"] == 1
    assert tool_result["cards"][0]["knowledgeType"] == "decision"

    access_token, _ = AuthService.create_access_token(test_user.id)
    res = await client.get(
        "/api/v1/knowledge/search",
        params={"q": "噪声"},
        headers={"Authorization": f"Bearer {access_token}"},
    )
    assert res.status_code == 200, res.text
    assert res.json()["total"] == 1


@pytest.mark.asyncio
async def test_note_sources_lists_origin_documents(db_session, test_user, client):
    """知识笔记可以追溯回源文档与卡片。"""
    from app.services.note.note_service import NoteService

    document, _duplicate = await ingest(
        db_session,
        test_user.id,
        source_type="upload",
        external_id="source-1",
        title="原始标准",
        text="原文。",
    )
    note = await NoteService.create_note(db_session, {
        "title": "知识笔记",
        "note_type": "markdown",
        "content": {"text": "正文"},
        "created_by": test_user.id,
        "updated_by": test_user.id,
    })
    db_session.add(NoteSource(note_id=note.id, document_id=document.id, role="primary"))
    await db_session.commit()

    access_token, _ = AuthService.create_access_token(test_user.id)
    res = await client.get(
        f"/api/v1/knowledge/notes/{note.id}/sources",
        headers={"Authorization": f"Bearer {access_token}"},
    )
    assert res.status_code == 200, res.text
    items = res.json()["items"]
    assert len(items) == 1
    assert items[0]["sourceType"] == "upload"
    assert items[0]["title"] == "原始标准"
