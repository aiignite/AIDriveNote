"""DOCX 解析：段落（保留标题层级）+ 表格 + 页眉页脚。"""
from __future__ import annotations

import io
import re

from app.services.knowledge.parsing.base import (
    ParsedDocument,
    ParseError,
    normalize_text,
    truncate_text,
)

_HEADING_RE = re.compile(r"heading\s*(\d+)", re.IGNORECASE)


def _heading_prefix(style_name: str) -> str:
    match = _HEADING_RE.search(style_name or "")
    if not match:
        return ""
    level = min(max(int(match.group(1)), 1), 6)
    return "#" * level + " "


def parse_docx(data: bytes) -> ParsedDocument:
    try:
        from docx import Document
    except Exception as exc:  # noqa: BLE001
        raise ParseError("python-docx 未安装，无法解析 DOCX") from exc

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"DOCX 读取失败: {exc}") from exc

    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style is not None else ""
        parts.append(f"{_heading_prefix(style_name)}{text}")

    for table in doc.tables:
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))

    for section in doc.sections:
        for container in (section.header, section.footer):
            if container is None:
                continue
            text = "\n".join(
                paragraph.text.strip()
                for paragraph in container.paragraphs
                if paragraph.text and paragraph.text.strip()
            )
            if text:
                parts.append(text)

    body = normalize_text("\n\n".join(parts))
    body, truncated = truncate_text(body)
    return ParsedDocument(
        text=body,
        parser="docx",
        meta={"truncated": truncated, "paragraphs": len(doc.paragraphs), "tables": len(doc.tables)},
    )
