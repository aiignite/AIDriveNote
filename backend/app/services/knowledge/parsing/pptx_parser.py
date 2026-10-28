"""PPTX 解析：按幻灯片提取文本框与表格。"""
from __future__ import annotations

import io

from app.services.knowledge.parsing.base import (
    ParsedDocument,
    ParseError,
    normalize_text,
    truncate_text,
)


def parse_pptx(data: bytes) -> ParsedDocument:
    try:
        from pptx import Presentation
    except Exception as exc:  # noqa: BLE001
        raise ParseError("python-pptx 未安装，无法解析 PPTX") from exc

    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PPTX 读取失败: {exc}") from exc

    parts: list[str] = []
    for index, slide in enumerate(presentation.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text or "").strip()
                if text:
                    texts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        texts.append(" | ".join(cells))
        if texts:
            parts.append(f"## Slide {index}\n" + "\n".join(texts))

    body = normalize_text("\n\n".join(parts))
    body, truncated = truncate_text(body)
    return ParsedDocument(
        text=body,
        parser="pptx",
        meta={"slides": len(presentation.slides), "truncated": truncated},
    )
