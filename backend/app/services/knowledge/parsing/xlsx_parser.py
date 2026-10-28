"""XLSX 解析：按工作表输出，限制页数/行数/字符数。"""
from __future__ import annotations

import io

from app.services.knowledge.parsing.base import (
    MAX_EXTRACTED_CHARS,
    ParsedDocument,
    ParseError,
    normalize_text,
    truncate_text,
)

MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 500


def parse_xlsx(data: bytes) -> ParsedDocument:
    try:
        from openpyxl import load_workbook
    except Exception as exc:  # noqa: BLE001
        raise ParseError("openpyxl 未安装，无法解析 XLSX") from exc

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Excel 读取失败: {exc}") from exc

    parts: list[str] = []
    total_chars = 0
    sheet_names = list(workbook.sheetnames)
    try:
        for index, sheet_name in enumerate(sheet_names):
            if index >= MAX_SHEETS:
                parts.append(f"...（另有 {len(sheet_names) - MAX_SHEETS} 个工作表未展示）")
                break
            sheet = workbook[sheet_name]
            lines: list[str] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_index > MAX_ROWS_PER_SHEET:
                    lines.append(f"...（仅展示前 {MAX_ROWS_PER_SHEET} 行）")
                    break
                cells = [
                    str(cell).strip()
                    for cell in row
                    if cell is not None and str(cell).strip()
                ]
                if cells:
                    lines.append(" | ".join(cells))
            if lines:
                block = f"## {sheet_name}\n" + "\n".join(lines)
                parts.append(block)
                total_chars += len(block)
                if total_chars >= MAX_EXTRACTED_CHARS:
                    break
    finally:
        workbook.close()

    body = normalize_text("\n\n".join(parts))
    body, truncated = truncate_text(body)
    return ParsedDocument(
        text=body,
        parser="xlsx",
        meta={
            "sheets": len(sheet_names),
            "truncated": truncated or total_chars >= MAX_EXTRACTED_CHARS,
        },
    )
