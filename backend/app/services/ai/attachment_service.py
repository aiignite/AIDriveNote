"""AI chat attachment storage and multimodal loading."""
from __future__ import annotations

import base64
import logging
import os
import uuid
from typing import Any

from fastapi import UploadFile
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.ai import AIAttachment

logger = logging.getLogger(__name__)

_BACKEND_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
_UPLOAD_DIR = os.path.join(_BACKEND_DIR, "uploads")
_MAX_UPLOAD_SIZE = 20 * 1024 * 1024
_MAX_VISION_IMAGE_BYTES = 10 * 1024 * 1024
_FALLBACK_MIME = "application/octet-stream"
_MAX_EXTRACT_CHARS = 50_000
_TRUNCATION_HINT = "\n\n... (内容过长，已截断)"
_MAX_XLSX_ROWS_PER_SHEET = 200
_MAX_XLSX_SHEETS = 20
_IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"})
_TEXT_EXTENSIONS = frozenset({".txt", ".md", ".csv", ".json", ".xml"})
_EXCEL_EXTENSIONS = frozenset({".xlsx", ".xlsm"})
_EXTENSION_MIME_MAP = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".json": "application/json",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
}
_ALLOWED_MIME_PREFIXES = (
    "image/",
    "text/",
    "application/pdf",
    "application/json",
    "application/vnd.openxmlformats",
    "application/msword",
)


def _normalize_upload_mime(file_name: str | None, mime_type: str | None) -> str:
    normalized = (mime_type or _FALLBACK_MIME).lower()
    if normalized not in {_FALLBACK_MIME, "application/octet-stream"}:
        return normalized
    file_ext = os.path.splitext(file_name or "")[1].lower()
    return _EXTENSION_MIME_MAP.get(file_ext, normalized)


def _is_excel_file(file_name: str, mime_type: str | None = None) -> bool:
    lower = file_name.lower()
    if lower.endswith(".xls") and not lower.endswith((".xlsx", ".xlsm")):
        return True
    if any(lower.endswith(ext) for ext in _EXCEL_EXTENSIONS):
        return True
    mime = (mime_type or "").lower()
    return "spreadsheet" in mime or mime.endswith("ms-excel")


def _is_image_attachment(attachment: AIAttachment) -> bool:
    mime_type = attachment.mime_type or _FALLBACK_MIME
    file_name = (attachment.original_name or "").lower()
    return mime_type.startswith("image/") or any(file_name.endswith(ext) for ext in _IMAGE_EXTENSIONS)


class AttachmentService:
    @staticmethod
    def resolve_file_path(attachment: AIAttachment) -> str | None:
        if not attachment.file_path:
            return None
        resolved_path = attachment.file_path
        if not os.path.isfile(resolved_path):
            fallback = os.path.join(_BACKEND_DIR, attachment.file_path)
            if os.path.isfile(fallback):
                resolved_path = fallback
            else:
                return None
        return resolved_path

    @staticmethod
    async def upload(
        db: AsyncSession,
        user_id: uuid.UUID,
        file: UploadFile,
        *,
        conversation_id: uuid.UUID | None = None,
        purpose: str = "chat",
    ) -> AIAttachment:
        mime = _normalize_upload_mime(file.filename, file.content_type)
        if not any(mime.startswith(prefix) for prefix in _ALLOWED_MIME_PREFIXES):
            raise ValueError(f"不支持的文件类型: {mime}")

        bucket = str(conversation_id) if conversation_id else "_no_conv"
        upload_dir = os.path.join(_UPLOAD_DIR, "ai", str(user_id), bucket)
        os.makedirs(upload_dir, exist_ok=True)

        file_ext = os.path.splitext(file.filename or "")[1]
        file_id = str(uuid.uuid4())
        file_name = f"{file_id}{file_ext}"
        file_path = os.path.join(upload_dir, file_name)

        content = await file.read()
        file_size = len(content)
        if file_size > _MAX_UPLOAD_SIZE:
            raise ValueError(
                f"文件过大 ({file_size // (1024 * 1024)}MB > {_MAX_UPLOAD_SIZE // (1024 * 1024)}MB)"
            )

        with open(file_path, "wb") as f:
            f.write(content)

        width: int | None = None
        height: int | None = None
        if mime.startswith("image/"):
            try:
                from PIL import Image  # type: ignore

                with Image.open(file_path) as img:
                    width, height = img.size
            except Exception:
                width = None
                height = None

        attachment = AIAttachment(
            user_id=user_id,
            conversation_id=conversation_id,
            file_name=file_name,
            original_name=file.filename or "unknown",
            file_size=file_size,
            mime_type=mime,
            file_path=file_path,
            purpose=purpose,
            width=width,
            height=height,
        )
        db.add(attachment)
        await db.commit()
        await db.refresh(attachment)
        return attachment

    @staticmethod
    async def get(
        db: AsyncSession,
        user_id: uuid.UUID,
        attachment_id: uuid.UUID,
    ) -> AIAttachment | None:
        result = await db.execute(
            select(AIAttachment).where(
                AIAttachment.id == attachment_id,
                AIAttachment.user_id == user_id,
                AIAttachment.is_deleted == False,  # noqa: E712
            )
        )
        return result.scalar_one_or_none()

    @staticmethod
    async def fetch_map(
        db: AsyncSession,
        user_id: uuid.UUID,
        attachment_ids: list[str],
    ) -> dict[str, AIAttachment]:
        if not attachment_ids:
            return {}
        try:
            valid_uuids = [uuid.UUID(aid.strip()) for aid in attachment_ids if aid.strip()]
        except ValueError:
            return {}
        if not valid_uuids:
            return {}

        result = await db.execute(
            select(AIAttachment).where(
                AIAttachment.id.in_(valid_uuids),
                AIAttachment.user_id == user_id,
                AIAttachment.is_deleted == False,  # noqa: E712
            )
        )
        return {str(item.id): item for item in result.scalars().all()}

    @staticmethod
    async def resolve_names(
        db: AsyncSession,
        user_id: uuid.UUID,
        attachment_ids: list[str],
    ) -> list[str]:
        att_map = await AttachmentService.fetch_map(db, user_id, attachment_ids)
        return [
            att_map[aid].original_name
            for aid in attachment_ids
            if aid in att_map and att_map[aid].original_name
        ]

    @staticmethod
    async def partition_ids(
        db: AsyncSession,
        user_id: uuid.UUID,
        attachment_ids: list[str],
    ) -> tuple[list[str], list[str]]:
        att_map = await AttachmentService.fetch_map(db, user_id, attachment_ids)
        image_ids: list[str] = []
        document_ids: list[str] = []
        for attachment_id in attachment_ids:
            attachment = att_map.get(attachment_id.strip())
            if attachment and _is_image_attachment(attachment):
                image_ids.append(attachment_id)
            else:
                document_ids.append(attachment_id)
        return image_ids, document_ids

    @staticmethod
    async def load_images(
        db: AsyncSession,
        user_id: uuid.UUID,
        attachment_ids: list[str],
    ) -> tuple[list[dict[str, Any]], list[str]]:
        if not attachment_ids:
            return [], []

        att_map = await AttachmentService.fetch_map(db, user_id, attachment_ids)
        image_parts: list[dict[str, Any]] = []
        loaded_ids: list[str] = []

        for attachment_id in attachment_ids:
            attachment = att_map.get(attachment_id.strip())
            if not attachment or not _is_image_attachment(attachment):
                continue

            mime_type = attachment.mime_type or _FALLBACK_MIME
            file_name = (attachment.original_name or "").lower()
            resolved_path = AttachmentService.resolve_file_path(attachment)
            if not resolved_path:
                continue

            try:
                with open(resolved_path, "rb") as f:
                    file_bytes = f.read()
            except OSError as exc:
                logger.error("视觉加载图片失败 %s: %s", attachment_id, exc)
                continue

            if len(file_bytes) > _MAX_VISION_IMAGE_BYTES:
                logger.warning("图片过大，跳过视觉注入: %s", attachment.original_name)
                continue

            if not mime_type.startswith("image/"):
                for ext, mapped in _EXTENSION_MIME_MAP.items():
                    if file_name.endswith(ext) and mapped.startswith("image/"):
                        mime_type = mapped
                        break
                if not mime_type.startswith("image/"):
                    mime_type = "image/png"

            image_parts.append({
                "type": "image",
                "mimeType": mime_type,
                "data": base64.b64encode(file_bytes).decode("utf-8"),
            })
            loaded_ids.append(attachment_id)

        return image_parts, loaded_ids

    @staticmethod
    async def serialize_for_messages(
        db: AsyncSession,
        user_id: uuid.UUID,
        attachment_ids: list[str],
    ) -> list[dict[str, str]]:
        if not attachment_ids:
            return []
        att_map = await AttachmentService.fetch_map(db, user_id, attachment_ids)
        items: list[dict[str, str]] = []
        for attachment_id in attachment_ids:
            attachment = att_map.get(attachment_id.strip())
            if not attachment:
                continue
            items.append({
                "id": str(attachment.id),
                "name": attachment.original_name,
                "mimeType": attachment.mime_type,
            })
        return items

    @staticmethod
    def append_attachment_markers(message: str, file_names: list[str]) -> str:
        if not file_names:
            return message
        markers = "\n".join(f"📎 {name}" for name in file_names)
        return f"{message}\n\n{markers}".strip() if message.strip() else markers

    @staticmethod
    def build_document_hint(
        file_entries: list[tuple[str, str]],
        *,
        extraction_failed: bool = False,
    ) -> str:
        if not file_entries:
            return ""
        lines = "\n".join(f"  - {name} (attachment_id: {aid})" for aid, name in file_entries)
        if extraction_failed:
            return (
                f"\n\n[已上传文档附件但未能自动提取正文:\n{lines}\n"
                "请告知用户附件已收到，但当前无法解析该文件格式。]"
            )
        return f"\n\n[已上传文档附件:\n{lines}\n请结合下方附件正文回答。]"

    @staticmethod
    def _extract_pdf_content(file_path: str) -> str | None:
        try:
            from pypdf import PdfReader

            reader = PdfReader(file_path)
            parts: list[str] = []
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(f"--- 第 {i + 1} 页 ---\n{text.strip()}")
            if not parts:
                return "[PDF 未能提取到文本内容，可能是扫描件或图片型 PDF]"
            full_content = "\n\n".join(parts)
            if len(full_content) > _MAX_EXTRACT_CHARS:
                return full_content[:_MAX_EXTRACT_CHARS] + _TRUNCATION_HINT
            return full_content
        except ImportError:
            return "[缺少 PDF 处理库，请安装: pip install pypdf]"
        except Exception as exc:
            logger.error("PDF 提取失败: %s", exc)
            return f"[PDF 读取失败: {exc}]"

    @staticmethod
    def _extract_xlsx_content(file_path: str, *, original_name: str = "") -> str | None:
        file_lower = (original_name or file_path).lower()
        if file_lower.endswith(".xls") and not file_lower.endswith((".xlsx", ".xlsm")):
            return "[旧版 .xls 格式暂不支持，请另存为 .xlsx 后重新上传]"

        try:
            from openpyxl import load_workbook
        except ImportError:
            return "[缺少 Excel 处理库，请安装: pip install openpyxl]"

        try:
            wb = load_workbook(file_path, read_only=True, data_only=True)
            parts: list[str] = []
            total_chars = 0
            sheet_names = wb.sheetnames

            for sheet_idx, sheet_name in enumerate(sheet_names):
                if sheet_idx >= _MAX_XLSX_SHEETS:
                    remaining = len(sheet_names) - _MAX_XLSX_SHEETS
                    parts.append(f"... (另有 {remaining} 个工作表未展示)")
                    break

                ws = wb[sheet_name]
                sheet_lines: list[str] = []
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                    if row_idx > _MAX_XLSX_ROWS_PER_SHEET:
                        sheet_lines.append(
                            f"... (工作表「{sheet_name}」仅展示前 {_MAX_XLSX_ROWS_PER_SHEET} 行)"
                        )
                        break
                    cells = [
                        str(cell).strip()
                        for cell in row
                        if cell is not None and str(cell).strip()
                    ]
                    if cells:
                        sheet_lines.append(" | ".join(cells))

                if sheet_lines:
                    block = f"--- 工作表: {sheet_name} ---\n" + "\n".join(sheet_lines)
                    parts.append(block)
                    total_chars += len(block)
                    if total_chars > _MAX_EXTRACT_CHARS:
                        break

            wb.close()

            if not parts:
                return "[Excel 文件为空或未能读取到单元格数据]"

            full_content = "\n\n".join(parts)
            if len(full_content) > _MAX_EXTRACT_CHARS:
                return full_content[:_MAX_EXTRACT_CHARS] + _TRUNCATION_HINT
            return full_content
        except Exception as exc:
            logger.error("Excel 提取失败: %s", exc)
            return f"[Excel 读取失败: {exc}]"

    @staticmethod
    def _extract_file_content(attachment: AIAttachment, file_path: str) -> str | None:
        file_name = (attachment.original_name or "").lower()

        if file_name.endswith(".pdf") or (attachment.mime_type or "").endswith("pdf"):
            return AttachmentService._extract_pdf_content(file_path)

        if _is_excel_file(file_name, attachment.mime_type):
            return AttachmentService._extract_xlsx_content(
                file_path,
                original_name=attachment.original_name or "",
            )

        if any(file_name.endswith(ext) for ext in _TEXT_EXTENSIONS):
            for encoding in ("utf-8", "gbk", "latin-1"):
                try:
                    with open(file_path, encoding=encoding) as f:
                        content = f.read()
                    if len(content) > _MAX_EXTRACT_CHARS:
                        content = content[:_MAX_EXTRACT_CHARS] + _TRUNCATION_HINT
                    return content
                except UnicodeDecodeError:
                    continue
            return "[文本文件解码失败]"

        if file_name.endswith(".docx"):
            try:
                from docx import Document

                doc = Document(file_path)
                parts = [p.text for p in doc.paragraphs if p.text.strip()]
                full_content = "\n".join(parts)
                if len(full_content) > _MAX_EXTRACT_CHARS:
                    full_content = full_content[:_MAX_EXTRACT_CHARS] + _TRUNCATION_HINT
                return full_content or "[DOCX 文件为空]"
            except ImportError:
                return "[缺少 DOCX 处理库，请安装: pip install python-docx]"
            except Exception as exc:
                logger.error("DOCX 提取失败: %s", exc)
                return f"[DOCX 读取失败: {exc}]"

        return f"[暂不支持的文件类型: {attachment.original_name}]"

    @staticmethod
    async def extract_documents_text(
        db: AsyncSession,
        user_id: uuid.UUID,
        document_ids: list[str],
    ) -> str:
        if not document_ids:
            return ""

        att_map = await AttachmentService.fetch_map(db, user_id, document_ids)
        sections: list[str] = []

        for attachment_id in document_ids:
            attachment = att_map.get(attachment_id.strip())
            if not attachment or _is_image_attachment(attachment):
                continue
            resolved_path = AttachmentService.resolve_file_path(attachment)
            if not resolved_path:
                sections.append(f"## 文件: {attachment.original_name}\n[附件文件不存在或已被删除]")
                continue
            content = AttachmentService._extract_file_content(attachment, resolved_path)
            if content:
                sections.append(f"## 文件: {attachment.original_name}\n{content}")

        return "\n\n".join(sections)
